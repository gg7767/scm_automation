from django.contrib.auth.models import Group
from django.core.mail import EmailMessage
from django.core.management.base import BaseCommand

MONTHLY_PACK_ATTACHMENT_NAME = "scm_monthly_pack.xlsx"


def build_monthly_pack_workbook():
    """One workbook, one sheet per report in docs/PHASE4_INVENTORY_TRANSPORT_MACHINERY.md's
    reporting suite — the numbers here must match what's on-screen."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    from assets.services import machinery_utilization
    from billing.services import payables_aging, spend_analysis
    from logistics.services import freight_by_site_transporter_month
    from purchase.services import procurement_lead_time
    from stores.services import consumption_by_site_month, slow_moving_items

    def header(ws, cols):
        ws.append(cols)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

    wb = Workbook()

    ws = wb.active
    ws.title = "Spend by vendor"
    header(ws, ["Vendor", "Total spend"])
    for row in spend_analysis(group_by="vendor"):
        ws.append([row["label"], float(row["total"])])

    ws = wb.create_sheet("Spend by month")
    header(ws, ["Month", "Total spend"])
    for row in spend_analysis(group_by="month"):
        ws.append([row["label"], float(row["total"])])

    ws = wb.create_sheet("Lead time")
    header(ws, ["PO", "Vendor", "Indent to sent (days)", "Sent to 1st GRN (days)", "Sent to complete (days)"])
    for row in procurement_lead_time()["rows"]:
        ws.append([
            row["po"].po_number, row["vendor"], row["indent_to_sent_days"],
            row["sent_to_first_grn_days"], row["sent_to_complete_days"] or "",
        ])

    ws = wb.create_sheet("Payables aging")
    aging = payables_aging()
    header(ws, ["Bucket", "Amount"])
    for bucket, amount in aging["buckets"].items():
        ws.append([bucket, float(amount)])

    ws = wb.create_sheet("Slow-moving stock")
    header(ws, ["Site", "Item", "Quantity"])
    for balance in slow_moving_items():
        ws.append([balance.site.code, balance.item.name, float(balance.quantity)])

    ws = wb.create_sheet("Stock consumption")
    header(ws, ["Site", "Month", "Qty issued"])
    for row in consumption_by_site_month():
        ws.append([row["site"], row["month"], float(row["qty_issued"])])

    ws = wb.create_sheet("Freight")
    header(ws, ["Site", "Transporter", "Month", "Total freight"])
    for row in freight_by_site_transporter_month():
        ws.append([row["site"], row["transporter"], row["month"], float(row["total_freight"])])

    ws = wb.create_sheet("Machinery utilisation")
    header(ws, ["Machine", "Hours run", "Fuel (L)", "Fuel/hour", "Hire cost"])
    for row in machinery_utilization():
        ws.append([
            row["machine"].code, float(row["hours_run"]), float(row["fuel_litres"]),
            float(row["fuel_per_hour"]) if row["fuel_per_hour"] is not None else "",
            float(row["hire_cost"]) if row["hire_cost"] is not None else "",
        ])

    return wb


class Command(BaseCommand):
    help = (
        "Builds the monthly SCM pack (one Excel workbook, one sheet per report) "
        "and emails it to every SCM Head. Run monthly via cron, or on demand."
    )

    def handle(self, *args, **options):
        import io

        wb = build_monthly_pack_workbook()
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        recipients = list(
            Group.objects.filter(name="SCM Head").first().user_set.filter(is_active=True).values_list("email", flat=True)
        ) if Group.objects.filter(name="SCM Head").exists() else []
        recipients = [e for e in recipients if e]

        if not recipients:
            self.stdout.write(self.style.WARNING("No SCM Head with an email address found; workbook built but not sent."))
            return

        email = EmailMessage(
            subject="Monthly SCM Pack",
            body="Attached: this month's spend, lead-time, aging, stock, freight, and machinery reports.",
            to=recipients,
        )
        email.attach(
            MONTHLY_PACK_ATTACHMENT_NAME, buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        email.send(fail_silently=False)
        self.stdout.write(self.style.SUCCESS(f"Monthly SCM pack sent to {len(recipients)} recipient(s)."))
