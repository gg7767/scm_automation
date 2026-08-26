from openpyxl import Workbook
from openpyxl.styles import Font


def _header_row(ws, headers):
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)


def vendor_ledger_workbook(vendor, entries):
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendor Ledger"
    ws.append([f"Vendor: {vendor.code} — {vendor.name}"])
    _header_row(ws, ["Date", "Type", "Reference", "Amount", "Running balance"])
    for row in entries:
        ws.append([
            row["date"].strftime("%d-%m-%Y") if row["date"] else "",
            row["type"], row["ref"], float(row["amount"]), float(row["running_balance"]),
        ])
    return wb


def payables_aging_workbook(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Payables Aging"
    _header_row(ws, ["Bucket", "Amount"])
    for bucket, amount in data["buckets"].items():
        ws.append([bucket, float(amount)])

    detail_ws = wb.create_sheet("Bill Detail")
    _header_row(detail_ws, ["Bill Number", "Vendor", "Site", "Due Date", "Days Overdue", "Bucket", "Balance"])
    for row in data["rows"]:
        bill = row["bill"]
        detail_ws.append([
            bill.bill_number, bill.vendor.name, bill.site.code,
            bill.due_date.strftime("%d-%m-%Y") if bill.due_date else "",
            row["days_overdue"], row["bucket"], float(row["balance"]),
        ])
    return wb
